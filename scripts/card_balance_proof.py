#!/usr/bin/env python3
"""
Does the live rule actually even the cards out? Takes the survey as it stands
in Convex right now, plays the remaining interviews with the very rule the
server uses (`pickLeastUsed`), and writes the answer to an Excel workbook.

    python3 scripts/card_balance_proof.py acbus-demo
    python3 scripts/card_balance_proof.py acbus-demo 500 --runs 200

Nothing is written back to the database; this only reads.
"""
import json
import random
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def table(name):
    with tempfile.TemporaryFile() as output:
        subprocess.run(
            ['npx', 'convex', 'data', name, '--limit', '8000', '--format', 'jsonl'],
            stdout=output, stderr=subprocess.PIPE, check=True,
        )
        output.seek(0)
        return [json.loads(l) for l in output.read().decode('utf-8').splitlines() if l.startswith('{')]


def pick_least_used(sets, count, used):
    """The server's rule, line for line (see pickLeastUsed in convex/cardBalance.ts)."""
    pool = list(sets)
    random.shuffle(pool)               # ties come out in random order
    pool.sort(key=lambda s: used[s])   # stable: least-used first
    chosen = pool[:min(count, len(pool))]
    random.shuffle(chosen)             # and the slots they land in are random too
    return chosen


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    runs = 200 if '--runs' in sys.argv else 1
    if '--runs' in sys.argv:
        runs = int(sys.argv[sys.argv.index('--runs') + 1])
    if not args:
        sys.exit('usage: python3 scripts/card_balance_proof.py <survey id> [target] [--runs N]')
    survey = args[0]

    found = [q for q in table('questionnaires') if q['id'] == survey]
    if not found:
        sys.exit(f'No survey with id {survey}')
    questionnaire = found[0]
    target = int(args[1]) if len(args) > 1 else questionnaire['responseTarget']
    blocks = [q for q in questionnaire['questions'] if q['type'] == 'choice_experiment' and q['cards']]
    responses = [r for r in table('responses') if r['questionnaireId'] == survey]

    # Where each block stands today, straight from the recorded answers.
    today, with_cards = [], 0
    for block in blocks:
        counts = Counter({card['set']: 0 for card in block['cards']})
        for response in responses:
            for scenario in ((response.get('answers') or {}).get(block['id']) or {}).get('scenarios', []):
                if isinstance(scenario.get('set'), int):
                    counts[scenario['set']] += 1
        today.append(counts)
    with_cards = sum(
        1 for r in responses
        if any(((r.get('answers') or {}).get(b['id']) or {}).get('scenarios') for b in blocks)
    )
    remaining = max(0, target - with_cards)

    # The cards already shown, one row per recorded response, in survey order.
    dealt = [[] for _ in blocks]
    for response in sorted(responses, key=lambda r: r['serial']):
        shown = [
            [s['set'] for s in ((response.get('answers') or {}).get(b['id']) or {}).get('scenarios', [])
             if isinstance(s.get('set'), int)]
            for b in blocks
        ]
        if not any(shown):
            continue
        for row, block, seen in zip(dealt, blocks, shown):
            row.append((seen + [''] * block['scenariosPerRespondent'])[:block['scenariosPerRespondent']])

    # Play the remaining interviews, many times over, with the server's rule.
    spreads, finals, trace = [], [], []
    for run in range(runs):
        counts = [Counter(c) for c in today]
        for n in range(remaining):
            for i, (block, used) in enumerate(zip(blocks, counts)):
                drawn = pick_least_used(list(used), block['scenariosPerRespondent'], used)
                for s in drawn:
                    used[s] += 1
                if run == 0:
                    dealt[i].append(drawn)
            if run == 0:
                trace.append([with_cards + n + 1] + [max(u.values()) - min(u.values()) for u in counts])
        spreads.append([max(u.values()) - min(u.values()) for u in counts])
        finals.append(counts)
    worst = [max(s[i] for s in spreads) for i in range(len(blocks))]
    final = finals[0]

    book = Workbook()
    head = Font(bold=True)
    good = PatternFill('solid', fgColor='E1F5EE')
    bad = PatternFill('solid', fgColor='FCEBEB')
    thin = Side(style='thin', color='9A9A9A')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    middle = Alignment(horizontal='center', vertical='center')
    FILL = {
        'set': PatternFill('solid', fgColor='BDD7EE'),
        'card': PatternFill('solid', fgColor='C4CE8F'),
        'option': PatternFill('solid', fgColor='CCD0EC'),
        'cardno': PatternFill('solid', fgColor='DDE5B6'),
        'freq': PatternFill('solid', fgColor='E3B7B8'),
        'recorded': PatternFill('solid', fgColor='FFF2CC'),
    }

    def banner(sheet, cells, value, fill):
        sheet.merge_cells(cells)
        first = sheet[cells.split(':')[0]]
        first.value = value
        first.font = Font(bold=True)
        first.alignment = middle
        for row in sheet[cells]:
            for cell in row:
                cell.fill = FILL[fill]
                cell.border = box

    for n, (block, now, end_counts) in enumerate(zip(blocks, today, final), 1):
        k = block['scenariosPerRespondent']
        label = (block['label']['en'] or block['label']['bn'] or f'Block {n}').strip()
        rows = dealt[n - 1]
        sheet = book.create_sheet(f'Block {n}')

        # Left: one row per respondent, one column per option.
        last = get_column_letter(1 + k)
        banner(sheet, f'A1:A3', 'Set', 'set')
        banner(sheet, f'B1:{last}1', 'Card', 'card')
        banner(sheet, f'B2:{last}2', label, 'card')
        for i in range(k):
            cell = sheet.cell(3, 2 + i, f'Option{i + 1}')
            cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), middle, FILL['option'], box

        # Right: how often each card falls in each option column.
        gap = 2 + k                      # one blank column between the tables
        card_col = gap + 1
        first_freq = card_col + 1
        card_letter = get_column_letter(card_col)
        freq_first = get_column_letter(first_freq)
        freq_last = get_column_letter(first_freq + k)
        banner(sheet, f'{card_letter}1:{card_letter}3', 'Card', 'cardno')
        banner(sheet, f'{freq_first}1:{freq_last}1', 'Frequency', 'freq')
        banner(sheet, f'{freq_first}2:{freq_last}2', label, 'freq')
        for i in range(k + 1):
            cell = sheet.cell(3, first_freq + i, 'Total' if i == k else f'Option{i + 1}')
            cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), middle, FILL['option'], box

        for r, cards in enumerate(rows, start=4):
            set_cell = sheet.cell(r, 1, r - 3)
            set_cell.border = box
            if r - 4 < with_cards:
                set_cell.fill = FILL['recorded']   # actually collected, not simulated
            for i, card in enumerate(cards):
                sheet.cell(r, 2 + i, card).border = box

        per_option = [Counter() for _ in range(k)]
        for cards in rows:
            for i, card in enumerate(cards):
                if card != '':
                    per_option[i][card] += 1
        for r, card in enumerate(sorted(now), start=4):
            sheet.cell(r, card_col, card).border = box
            for i in range(k):
                sheet.cell(r, first_freq + i, per_option[i][card]).border = box
            total = sheet.cell(r, first_freq + k, sum(p[card] for p in per_option))
            total.border, total.font = box, Font(bold=True)

        sheet.column_dimensions['A'].width = 7
        for i in range(k):
            sheet.column_dimensions[get_column_letter(2 + i)].width = 9
        sheet.column_dimensions[get_column_letter(gap)].width = 3
        sheet.column_dimensions[card_letter].width = 7
        for i in range(k + 1):
            sheet.column_dimensions[get_column_letter(first_freq + i)].width = 9
        sheet.freeze_panes = 'A4'

    summary = book['Sheet']
    summary.title = 'Summary'
    book.move_sheet('Summary', offset=-len(blocks))
    sheet = summary
    sheet.append([f'Card balance for {survey}'])
    sheet['A1'].font = Font(bold=True, size=14)
    sheet.append([])
    for label, value in [
        ('Survey target', target),
        ('Responses with cards so far', with_cards),
        ('Interviews simulated per run', remaining),
        ('Runs simulated', runs),
        ('Rule used', 'least-used cards first, ties at random (the server rule)'),
    ]:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = head
    sheet.append([])
    sheet.append(['Block', 'Cards', 'Per respondent', 'Shown today',
                  'Gap today', f'Gap at {target}', f'Worst gap in {runs} runs', 'Verdict'])
    for cell in sheet[sheet.max_row]:
        cell.font = head
    for n, (block, now, end_counts) in enumerate(zip(blocks, today, final), 1):
        gap_now = max(now.values()) - min(now.values())
        gap_end = max(end_counts.values()) - min(end_counts.values())
        sheet.append([
            f'Block {n}', len(block['cards']), block['scenariosPerRespondent'],
            sum(now.values()), gap_now, gap_end, worst[n - 1],
            'within 1' if worst[n - 1] <= 1 else 'FAILS',
        ])
        sheet.cell(sheet.max_row, 8).fill = good if worst[n - 1] <= 1 else bad
    for i, width in enumerate([28, 10, 16, 14, 12, 14, 22, 12], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    sheet = book.create_sheet('Closing the gap')
    sheet.append(['After response'] + [f'Block {n} gap' for n in range(1, len(blocks) + 1)])
    for cell in sheet[1]:
        cell.font = head
    for row in trace:
        sheet.append(row)
    sheet.freeze_panes = 'A2'

    sheet = book.create_sheet('How to read this')
    for line in [
        'What this workbook is',
        '',
        'One sheet per choice block. On the left, one row per respondent: the cards that',
        'respondent sees, one column per scenario slot (Option1, Option2, ...). On the right,',
        'how often each card falls in each slot, and its total across the block.',
        '',
        'Rows with a yellow Set number are responses already collected. The rest are played out',
        'with the same rule the server uses when a tablet starts an interview: hand out the cards',
        'used least so far, ties broken at random.',
        '',
        'The guarantee is on the Total column: every card shown equally often across the block,',
        'give or take one. The per-Option columns are near-even but not guaranteed, because a',
        'respondent\'s cards are shuffled into their slots at random.',
        '',
        'Gap = (most-shown card) minus (least-shown card) within a block. 0 or 1 is the goal.',
        '',
        'Nothing here is written back to the database, and the website does not read this file.',
    ]:
        sheet.append([line])
    sheet['A1'].font = Font(bold=True, size=12)
    sheet.column_dimensions['A'].width = 100

    out = Path('rough')
    out.mkdir(exist_ok=True)
    path = out / f'card_balance_proof_{survey}.xlsx'
    book.save(path)

    print(f'{survey}: target {target}, {with_cards} responses with cards, '
          f'{remaining} simulated x {runs} runs\n')
    for n, (block, now, end) in enumerate(zip(blocks, today, final), 1):
        print(f'Block {n}: {len(block["cards"])} cards, {block["scenariosPerRespondent"]} per respondent')
        print(f'  today       shown {sum(now.values()):>5}   least {min(now.values())}  most {max(now.values())}  gap {max(now.values()) - min(now.values())}')
        print(f'  at {target:<9} shown {sum(end.values()):>5}   least {min(end.values())}  most {max(end.values())}  gap {max(end.values()) - min(end.values())}')
        print(f'  worst gap across {runs} runs: {worst[n - 1]}')
    print(f'\nWrote {path}')


if __name__ == '__main__':
    main()
